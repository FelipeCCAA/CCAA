"""Contrato OP-07 por API y lectura real de los archivos generados."""
import csv
import json
import time
import tracemalloc
from datetime import date
from decimal import Decimal
from io import BytesIO, StringIO
from unittest.mock import patch

from openpyxl import load_workbook
from django.contrib.auth.models import User
from usuarios.models import PerfilUsuario

from recepcion.models import ModuloRecepcion, Recepcion
from recepcion.tests import BaseAPIRecepcion


class ExportacionesOP07Tests(BaseAPIRecepcion):
    ruta = "/api/recepcion/recepciones/resumen-diario/"

    def crear(self, **extra):
        datos = dict(fecha=date(2026, 7, 31), guia="0000123",
                     estado=Recepcion.Estado.REGISTRADA,
                     tipo_leche=Recepcion.TipoLeche.ENTERA,
                     litros=Decimal("1234.56"))
        datos.update(extra)
        return Recepcion.objects.create(**datos)

    def obtener(self, formato="json", **parametros):
        return self.cliente.get(self.ruta, dict(fecha="2026-07-31", formato=formato,
                                              detalle="1", **parametros))

    def test_csv_neutraliza_textos_formula_sin_alterar_numeros(self):
        peligrosos = ["=1+1", "+1+1", "-1+1", "@SUM(A1)", " \t=1+1", "\r=1+1"]
        for texto in peligrosos:
            self.crear(guia=texto)
        filas = list(csv.DictReader(StringIO(self.obtener("csv").content.decode("utf-8-sig")), delimiter=";"))
        self.assertEqual([fila["Guía"] for fila in filas], ["'" + x for x in peligrosos])
        self.assertEqual({fila["Litros"] for fila in filas}, {"1234.56"})

    def test_xlsx_conserva_texto_literal_sin_formulas(self):
        for guia in ["=1+1", "0000123", "+1+1", "\t=1+1"]:
            self.crear(guia=guia)
        libro = load_workbook(BytesIO(self.obtener("xlsx").content), data_only=False)
        self.assertEqual([libro.active.cell(fila, 3).value for fila in range(2, 6)],
                         ["=1+1", "0000123", "+1+1", "\t=1+1"])
        self.assertFalse(any(celda.data_type == "f" for hoja in libro for fila in hoja for celda in fila))

    def test_tres_formatos_coinciden_en_filtros_cantidades_y_faltantes(self):
        recepcion = self.crear()
        ModuloRecepcion.objects.create(recepcion=recepcion, numero=1, crioscopia=Decimal("-0.521"))
        self.crear(fecha=date(2026, 8, 1), litros=Decimal("9000"))
        self.crear(estado=Recepcion.Estado.BORRADOR)
        self.crear(estado=Recepcion.Estado.ANULADA)
        datos = self.obtener().data
        filas = list(csv.DictReader(StringIO(self.obtener("csv").content.decode("utf-8-sig")), delimiter=";"))
        libro = load_workbook(BytesIO(self.obtener("xlsx").content))
        totales = dict(libro["Totales"].values)
        self.assertEqual(datos["camiones"], 1)
        self.assertEqual(len(filas), 1)
        self.assertEqual(libro.active.max_row, 2)
        self.assertEqual(filas[0]["Litros"], datos["litros"])
        self.assertEqual(str(totales["Litros"]), datos["litros"])
        self.assertEqual(filas[0]["Kg romana"], "")
        self.assertIsNone(libro.active.cell(2, 9).value)
        self.assertEqual(totales["Camiones sin romana"], 1)
        self.assertEqual(filas[0]["Crioscopías"], "M1: -0.521")

    def test_limite_exacto_no_trunca_y_exceso_rechaza_los_tres_formatos(self):
        self.crear()
        self.crear()
        with patch("recepcion.views.MAX_RECEPCIONES_REPORTE", 2, create=True):
            self.assertEqual(len(self.obtener().data["detalle"]), 2)
            self.crear()
            for formato in ("json", "csv", "xlsx"):
                with self.subTest(formato=formato):
                    respuesta = self.obtener(formato)
                    self.assertEqual(respuesta.status_code, 400)
                    self.assertEqual(respuesta.data["code"], "REPORTE_DEMASIADO_GRANDE")
                    self.assertIn("reduce", respuesta.data["detail"].lower())
                    self.assertNotIn("Content-Disposition", respuesta)

    def test_exportacion_exige_autenticacion_y_permiso_del_modulo(self):
        self.crear()
        self.cliente.credentials()
        for formato in ("json", "csv", "xlsx"):
            self.assertIn(self.obtener(formato).status_code, (401, 403))
        usuario = User.objects.create_user("op07-aseo")
        PerfilUsuario.objects.create(usuario=usuario, area="aseo", nivel="trabajador", rol="aseo")
        self.cliente.force_authenticate(usuario)
        for formato in ("json", "csv", "xlsx"):
            self.assertEqual(self.obtener(formato).status_code, 403)

    def test_reporte_vacio_y_fechas_invalidas_sin_archivo(self):
        self.assertEqual(self.obtener().data["camiones"], 0)
        for parametros in ({"desde": "2026-08-01", "hasta": "2026-07-31"},
                           {"fecha": "2026-13-99"}):
            respuesta = self.cliente.get(self.ruta, {**parametros, "formato": "xlsx"})
            self.assertEqual(respuesta.status_code, 400)
            self.assertNotIn("Content-Disposition", respuesta)

    def test_tope_real_mil_recepciones_cuatro_modulos_y_exceso(self):
        # Fixtures masivas deliberadas: no representan operaciones de planta.
        # Medición local reproducible, sin afirmaciones de capacidad productiva.
        from recepcion.exportaciones import MAX_RECEPCIONES_REPORTE

        self.assertEqual(MAX_RECEPCIONES_REPORTE, 1000)
        filas = Recepcion.objects.bulk_create([
            Recepcion(fecha=date(2026, 7, 31), guia=f"{i:060d}",
                      estado=Recepcion.Estado.REGISTRADA,
                      tipo_leche=Recepcion.TipoLeche.ENTERA, litros=Decimal("1234.56"))
            for i in range(1000)
        ])
        ModuloRecepcion.objects.bulk_create([
            ModuloRecepcion(recepcion=r, numero=n, crioscopia=Decimal("-0.521"))
            for r in filas for n in range(1, 5)
        ])
        for formato in ("json", "csv", "xlsx"):
            tracemalloc.start()
            inicio = time.perf_counter()
            try:
                respuesta = self.obtener(formato)
                _, pico = tracemalloc.get_traced_memory()
            finally:
                tracemalloc.stop()
            self.assertEqual(respuesta.status_code, 200)
            self.assertLess(pico, 64 * 1024 * 1024)
            print("OP07_REPORTE_BENCH " + json.dumps({
                "formato": formato, "recepciones": 1000, "modulos": 4000,
                "segundos": round(time.perf_counter() - inicio, 3), "pico_bytes_python": pico,
            }), flush=True)
            if formato == "json":
                self.assertEqual(len(respuesta.data["detalle"]), 1000)
                self.assertEqual(respuesta.data["litros"], "1234560.00")
            elif formato == "csv":
                self.assertEqual(len(list(csv.reader(StringIO(respuesta.content.decode("utf-8-sig")), delimiter=";"))), 1001)
            else:
                libro = load_workbook(BytesIO(respuesta.content), read_only=True)
                self.assertEqual(libro.active.max_row, 1001)
                libro.close()
        self.crear()
        for formato in ("json", "csv", "xlsx"):
            self.assertEqual(self.obtener(formato).status_code, 400)
